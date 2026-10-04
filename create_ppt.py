from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Create presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
PRIMARY = RGBColor(0x5B, 0x5F, 0xEF)
SECONDARY = RGBColor(0x00, 0xC2, 0xFF)
DARK = RGBColor(0x1C, 0x1C, 0x1E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF2, 0xF2, 0xF7)

def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    
    # Background shape
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK
    shape.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11), Inches(1))
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(24)
    p.font.color.rgb = SECONDARY
    p.alignment = PP_ALIGN.CENTER
    
    return slide

def add_content_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = DARK
    
    # Content
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(12), Inches(5.5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, bullet in enumerate(bullets):
        if i > 0:
            tf.add_paragraph()
        p = tf.paragraphs[i]
        p.text = bullet
        p.font.size = Pt(20)
        p.font.color.rgb = DARK
        p.space_after = Pt(12)
        p.level = 0
    
    return slide

def add_two_column_slide(prs, title, left_title, left_bullets, right_title, right_bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = DARK
    
    # Left column title
    left_title_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(5.5), Inches(0.5))
    tf = left_title_box.text_frame
    p = tf.paragraphs[0]
    p.text = left_title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    
    # Left column content
    left_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.9), Inches(5.5), Inches(5))
    tf = left_box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(left_bullets):
        if i > 0:
            tf.add_paragraph()
        p = tf.paragraphs[i]
        p.text = bullet
        p.font.size = Pt(16)
        p.font.color.rgb = DARK
        p.space_after = Pt(8)
    
    # Right column title
    right_title_box = slide.shapes.add_textbox(Inches(7), Inches(1.3), Inches(5.5), Inches(0.5))
    tf = right_title_box.text_frame
    p = tf.paragraphs[0]
    p.text = right_title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = SECONDARY
    
    # Right column content
    right_box = slide.shapes.add_textbox(Inches(7), Inches(1.9), Inches(5.5), Inches(5))
    tf = right_box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(right_bullets):
        if i > 0:
            tf.add_paragraph()
        p = tf.paragraphs[i]
        p.text = bullet
        p.font.size = Pt(16)
        p.font.color.rgb = DARK
        p.space_after = Pt(8)
    
    return slide

def add_stats_slide(prs, title, stats):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = DARK
    
    # Stats
    for i, (number, label) in enumerate(stats):
        left = Inches(1 + i * 3)
        
        # Number
        num_box = slide.shapes.add_textbox(left, Inches(2), Inches(2.5), Inches(1))
        tf = num_box.text_frame
        p = tf.paragraphs[0]
        p.text = number
        p.font.size = Pt(48)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        p.alignment = PP_ALIGN.CENTER
        
        # Label
        label_box = slide.shapes.add_textbox(left, Inches(3.2), Inches(2.5), Inches(0.5))
        tf = label_box.text_frame
        p = tf.paragraphs[0]
        p.text = label
        p.font.size = Pt(18)
        p.font.color.rgb = DARK
        p.alignment = PP_ALIGN.CENTER
    
    return slide

# Slide 1: Title
add_title_slide(prs, "SkillVerse AI", "AI-Powered Career Preparation Platform")

# Slide 2: Project Overview
add_content_slide(prs, "Project Overview", [
    "• AI-powered placement and career preparation platform for students",
    "• Combines authentication, resume analysis, ATS scoring, job matching, and interview practice",
    "• Built with Django backend and React frontend",
    "• Features JWT-based authentication and modern UI/UX",
    "• Helps students become job-ready from one platform"
])

# Slide 3: Key Features
add_two_column_slide(prs, "Key Features",
    "Resume Analysis", [
        "• PDF, DOCX, TXT support",
        "• Instant ATS scoring",
        "• Section-by-section analysis",
        "• Keyword optimization tips",
        "• PDF report generation"
    ],
    "Interview Practice", [
        "• Role-specific questions",
        "• Instant AI evaluation",
        "• Detailed feedback",
        "• Performance tracking",
        "• Progress analytics"
    ]
)

# Slide 4: Technology Stack
add_two_column_slide(prs, "Technology Stack",
    "Frontend", [
        "• React 19 + TypeScript",
        "• Vite build tool",
        "• React Router",
        "• Axios for API calls",
        "• Framer Motion animations",
        "• Tailwind CSS"
    ],
    "Backend", [
        "• Python 3.11+",
        "• Django 5.2 + DRF",
        "• Simple JWT auth",
        "• SQLite database",
        "• PyMuPDF + python-docx",
        "• OpenAI SDK (optional)"
    ]
)

# Slide 5: Architecture
add_content_slide(prs, "System Architecture", [
    "• React Frontend → HTTP/REST API → Django Backend",
    "• Django REST Framework handles API requests",
    "• Three main apps: Accounts, Resume API, Interviews",
    "• SQLite database for development",
    "• Optional OpenAI integration for AI evaluation"
])

# Slide 6: API Endpoints
add_two_column_slide(prs, "API Endpoints",
    "Authentication", [
        "• POST /auth/register",
        "• POST /auth/login",
        "• POST /auth/token",
        "• GET /user/me"
    ],
    "Resume & Interview", [
        "• POST /resume/analyze",
        "• GET /interview/start",
        "• POST /interview/evaluate",
        "• POST /interview/summary",
        "• GET /interview/history",
        "• GET /interview/stats"
    ]
)

# Slide 7: ATS Scoring
add_content_slide(prs, "ATS Scoring System", [
    "• Maximum 100 points across 6 categories",
    "• Contact Info: 15 points (Name, Email, Phone)",
    "• Skills: 20 points (2 points per skill, max 10)",
    "• Education: 15 points",
    "• Projects: 20 points",
    "• Experience: 20 points",
    "• Links: 10 points (LinkedIn, GitHub)"
])

# Slide 8: Target Roles
add_content_slide(prs, "Supported Job Roles", [
    "• Python Developer - Python, FastAPI, Flask, Django, SQL, Git, Docker, AWS",
    "• Full Stack Developer - HTML, CSS, JavaScript, React, Node.js, Express, MongoDB, Git",
    "• Data Analyst - Python, SQL, Excel, Power BI, Pandas, NumPy",
    "• Java Developer - Java, Spring, SQL, Git, Docker",
    "• AI/ML Engineer - Python, Machine Learning, TensorFlow, Pandas, NumPy, Git"
])

# Slide 9: Statistics
add_stats_slide(prs, "Our Impact", [
    ("10K+", "Students"),
    ("500+", "Companies"),
    ("50K+", "Resumes Analyzed"),
    ("95%", "Success Rate")
])

# Slide 10: Roadmap
add_content_slide(prs, "Future Roadmap", [
    "• Full Resume Builder with templates",
    "• Aptitude Test Engine",
    "• Coding Challenge Engine",
    "• Skill Tracker & Certificate Manager",
    "• Company Preparation Roadmaps",
    "• AI Chatbot Mentor",
    "• Cover Letter Generator",
    "• Email Notifications & Leaderboards"
])

# Slide 11: Thank You
add_title_slide(prs, "Thank You", "Questions & Discussion")

# Save presentation
prs.save("C:\\Users\\THADURI SAI NAVEEN\\SkillVerse_AI\\code\\Skill_Verse_AI\\SkillVerse_AI_Presentation.pptx")
print("Presentation created successfully!")

